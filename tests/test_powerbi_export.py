"""Offline tests for the Power BI export, on synthetic prices."""
import numpy as np
import pandas as pd
import pytest

from src import config, data_loader, powerbi_export as pe


@pytest.fixture
def built(tmp_path, monkeypatch):
    rng = np.random.default_rng(1)
    funds = pe.load_fund_list()
    idx = pd.bdate_range("2017-09-01", "2023-06-30")
    raw = {}
    for _, f in funds.iterrows():
        p = 100 * np.exp(np.cumsum(rng.normal(0.0003, 0.01, len(idx))))
        if f["quote_currency_expected"] == "GBp":
            p *= 100
        if f["ticker"] == "V3AB.L":          # launched 2021
            p[idx < "2021-03-25"] = np.nan
        raw[f["ticker"]] = p
    raw = pd.DataFrame(raw, index=idx)
    fx = pd.DataFrame({"EURGBP": 0.88, "GBPUSD": 1.25}, index=idx)
    ccys = dict(zip(funds["ticker"], funds["quote_currency_expected"]))

    monkeypatch.setattr(pe, "PBI_DATA", tmp_path / "pbi")
    monkeypatch.setattr(config, "DATA_PROCESSED", tmp_path / "proc")
    monkeypatch.setattr(data_loader, "download_yahoo", lambda *a, **k: raw)
    monkeypatch.setattr(data_loader, "yahoo_currencies", lambda t: ccys)
    monkeypatch.setattr(data_loader, "download_fx", lambda refresh: fx)
    monkeypatch.setattr(data_loader, "download_sonia",
                        lambda refresh: pd.Series(1.0, index=idx))
    out = pe.build(verbose=False, end=pd.Timestamp("2023-05-31"))
    out["dir"] = tmp_path / "pbi"
    out["raw"] = raw
    return out


def test_files_written(built):
    for name in ["dim_fund", "dim_date", "dim_regime", "fact_returns_monthly",
                 "fact_sector_weights", "metrics_check"]:
        assert (built["dir"] / f"{name}.csv").exists()


def test_each_fund_keeps_its_own_start(built):
    d = built["dim_fund"].set_index("ticker")
    assert d.loc["SWDA.L", "FirstMonth"] == pd.Timestamp("2017-10-31")
    assert d.loc["V3AB.L", "FirstMonth"] == pd.Timestamp("2021-04-30")


def test_currency_conversion_is_neutral_with_flat_fx(built):
    # with flat FX, GBP returns must equal the raw local-currency returns
    f = built["fact"]
    swda = f[f["fund_key"] == 1].set_index("Date")["Return"]
    raw_m = built["raw"]["SWDA.L"].resample("ME").last().pct_change().loc[swda.index]
    assert np.allclose(swda, raw_m)


def test_benchmark_and_flags(built):
    f = built["fact"]
    esg = f[f["fund_key"] == 2]                      # SUSW vs SWDA
    swda = f[f["fund_key"] == 1].set_index("Date")["Return"]
    assert np.allclose(esg["BenchmarkReturn"], esg["Date"].map(swda))
    assert np.allclose(esg["ActiveReturn"], esg["Return"] - esg["BenchmarkReturn"])
    assert f.loc[f["fund_key"] == 1, "BenchmarkReturn"].isna().all()
    common = f.loc[f["in_common_period"] == 1, "Date"].min()
    assert common == pd.Timestamp("2021-04-30")


def test_regimes_cover_every_month(built):
    d = built["dim_date"]
    assert d["Regime"].notna().all()
    assert d.loc[d["Date"] == "2022-06-30", "Regime"].item() == "4. Energy & inflation shock"
    assert d.loc[d["Date"] == "2020-03-31", "Regime"].item() == "2. COVID crash"


def test_metric_formulas():
    # 3 months: +10%, -10%, +10%; rf 0
    dates = pd.date_range("2020-01-31", periods=3, freq="ME")
    fact = pd.DataFrame({"Date": dates, "fund_key": 1, "Return": [0.1, -0.1, 0.1],
                         "ExcessReturn": [0.1, -0.1, 0.1], "BenchmarkReturn": np.nan,
                         "ActiveReturn": np.nan, "RF": 0.0, "in_pair_period": 1})
    m = pe.metrics(fact).loc[1]
    assert m["CumulativeReturn"] == pytest.approx(1.1 * 0.9 * 1.1 - 1)
    assert m["MaxDrawdown"] == pytest.approx(-0.1)
    # a loss in the very first month counts as a drawdown from the £1 invested
    fact.loc[:, "Return"] = [-0.05, 0.02, 0.01]
    assert pe.metrics(fact).loc[1, "MaxDrawdown"] == pytest.approx(-0.05)
    assert m["CAGR"] == pytest.approx((1.089) ** 4 - 1)


def test_sector_template_not_overwritten(built, monkeypatch):
    path = built["dir"] / "fact_sector_weights.csv"
    df = pd.read_csv(path)
    df.loc[0, "Weight"] = 0.25
    df.to_csv(path, index=False)
    pe.build(verbose=False, end=pd.Timestamp("2023-05-31"))
    assert pd.read_csv(path).loc[0, "Weight"] == 0.25


def test_benchmark_metrics_equal_benchmark_fund_over_same_months(built):
    # SUSW's "benchmark" metrics must equal SWDA's own metrics over SUSW's months
    f = built["fact"]
    esg_dates = f.loc[f["fund_key"] == 2, "Date"]
    swda_same = f[(f["fund_key"] == 1) & f["Date"].isin(esg_dates)].copy()
    own = pe.metrics(swda_same).loc[1]
    m = built["metrics"].loc[2]
    assert m["BenchmarkCAGR"] == pytest.approx(own["CAGR"])
    assert m["BenchmarkSharpe"] == pytest.approx(own["Sharpe"])
    assert m["BenchmarkMaxDrawdown"] == pytest.approx(own["MaxDrawdown"])
    assert m["BenchmarkAnnVolatility"] == pytest.approx(own["AnnVolatility"])


def test_excel_model_round_trip(built):
    """The workbook for the web service must hold every table, real dates and true blanks."""
    from openpyxl import load_workbook
    path = built["dir"] / "ESG_funds_model.xlsx"
    wb = load_workbook(path)
    assert set(wb.sheetnames) == {"dim_fund", "dim_date", "dim_regime",
                                  "fact_returns_monthly", "fact_sector_weights"}
    for ws in wb.worksheets:
        assert list(ws.tables) == [ws.title]          # named Excel table per sheet
    x = pd.read_excel(path, sheet_name=None)
    fact = x["fact_returns_monthly"]
    assert len(fact) == len(built["fact"])
    assert pd.api.types.is_datetime64_any_dtype(fact["Date"])
    assert np.allclose(fact["Return"], built["fact"]["Return"])
    conv = fact[fact["fund_key"] == 1]
    assert conv["BenchmarkReturn"].isna().all()      # blank, not 0
    ws = wb["fact_returns_monthly"]
    hdr = [c.value for c in ws[1]]
    col = hdr.index("BenchmarkReturn") + 1
    assert ws.cell(row=2, column=col).value is None

"""
Build the Power BI dataset: a small star schema of CSV files.

Run from the repo root:
    python -m src.powerbi_export            # uses cached downloads if present
    python -m src.powerbi_export --refresh  # re-download everything

Writes to powerbi/data/:
    dim_fund.csv               one row per fund (from config/funds.csv + data dates)
    dim_date.csv               one row per month-end, with year, month and market regime
    dim_regime.csv             regime names, dates and sort order
    fact_returns_monthly.csv   one row per fund per month: return, risk-free, excess
                               return, benchmark return, active return, period flags
    fact_sector_weights.csv    TEMPLATE: fill in from each fund's factsheet
    metrics_check.csv          the headline metrics computed in Python, to check
                               that your DAX measures give the same numbers

Design choices:
  * Monthly data only. It is what the analysis uses, it keeps DAX fast, and the
    numbers match the factsheets' conventions.
  * Every price is converted to GBP (the experience of a UK investor).
  * Each fund keeps its own start date. Fair comparisons use the period flags:
      in_pair_period    months where the fund AND its benchmark both have data
      in_common_period  months where ALL funds have data
"""
from __future__ import annotations

import argparse
import warnings

import numpy as np
import pandas as pd

from . import cleaning, config, data_loader

FUNDS_CSV = config.ROOT / "config" / "funds.csv"
PBI_DATA = config.ROOT / "powerbi" / "data"
DOWNLOAD_START = "2017-09-01"  # Sep 2017 month-end price -> first return Oct 2017

REGIMES = [  # name, start, end (inclusive, month-end), short description
    # Names are numbered so they sort correctly in Power BI without "sort by column".
    ("1. Pre-COVID", "2000-01-31", "2020-01-31", "Steady growth, low rates"),
    ("2. COVID crash", "2020-02-29", "2020-04-30", "Global sell-off and rebound start"),
    ("3. Recovery", "2020-05-31", "2021-12-31", "Stimulus-driven rally, tech-led"),
    ("4. Energy & inflation shock", "2022-01-31", "2022-12-31",
     "War in Ukraine, energy prices spike, rates rise fast"),
    ("5. 2023 onwards", "2023-01-31", "2100-12-31", "Rate peak, AI-led rally"),
]

GICS_SECTORS = ["Information Technology", "Financials", "Health Care",
                "Consumer Discretionary", "Industrials", "Communication",
                "Consumer Staples", "Energy", "Materials", "Utilities",
                "Real Estate", "Cash and/or Derivatives"]


# --------------------------------------------------------------------------
# Inputs
# --------------------------------------------------------------------------
def load_fund_list(path=FUNDS_CSV) -> pd.DataFrame:
    funds = pd.read_csv(path, dtype={"benchmark_key": "Int64"})
    if funds["fund_key"].duplicated().any():
        raise ValueError("fund_key must be unique in funds.csv")
    bad = set(funds["benchmark_key"].dropna()) - set(funds["fund_key"])
    if bad:
        raise ValueError(f"benchmark_key refers to unknown fund_key(s): {bad}")
    return funds


def to_gbp(price: pd.Series, ccy: str, fx: pd.DataFrame,
           log: cleaning.QualityLog) -> pd.Series:
    """Convert a daily price series in `ccy` into GBP."""
    price, ccy = cleaning.to_major_units(price, ccy)
    price = cleaning.fix_unit_glitches(price, log)
    if ccy == "GBP":
        return price
    if ccy == "EUR":
        return cleaning.convert_to_gbp(price, fx["EURGBP"], log)
    if ccy == "USD":
        return cleaning.convert_to_gbp(price, 1 / fx["GBPUSD"], log)
    raise ValueError(f"{price.name}: no conversion rule for currency {ccy!r}")


def clean_daily(price: pd.Series, log: cleaning.QualityLog) -> pd.Series:
    """Weekdays only, from the fund's first price, short gaps carried forward."""
    s = price.dropna()
    s = s[s.index.dayofweek < 5]
    cal = pd.bdate_range(s.index.min(), s.index.max())
    s = s.reindex(cal)
    isna = s.isna()
    if isna.any():
        runs = isna.groupby((~isna).cumsum()).sum()
        log.add(f"{price.name}: {int(isna.sum())} missing weekday(s) "
                f"(includes UK holidays), longest run {int(runs.max())}")
    return s.ffill(limit=config.MAX_FFILL_DAYS)


# --------------------------------------------------------------------------
# Tables
# --------------------------------------------------------------------------
def regime_for(dates: pd.DatetimeIndex) -> pd.Series:
    out = pd.Series(pd.NA, index=dates, dtype="object")
    for name, start, end, _ in REGIMES:
        out[(dates >= start) & (dates <= end)] = name
    return out


def build_dim_date(months: pd.DatetimeIndex) -> pd.DataFrame:
    d = pd.DataFrame({"Date": months})
    d["Year"] = d["Date"].dt.year
    d["MonthNum"] = d["Date"].dt.month
    d["MonthName"] = d["Date"].dt.strftime("%b")
    d["YearMonth"] = d["Date"].dt.strftime("%Y-%m")
    d["Quarter"] = "Q" + d["Date"].dt.quarter.astype(str)
    d["Regime"] = regime_for(pd.DatetimeIndex(d["Date"])).values
    order = {r[0]: i + 1 for i, r in enumerate(REGIMES)}
    d["RegimeOrder"] = d["Regime"].map(order)
    return d


def build_dim_regime() -> pd.DataFrame:
    rows = [{"Regime": n, "RegimeOrder": i + 1,
             "StartMonth": None if s.startswith("2000") else s,
             "EndMonth": None if e.startswith("2100") else e, "Description": desc}
            for i, (n, s, e, desc) in enumerate(REGIMES)]
    return pd.DataFrame(rows)


def build_fact_returns(monthly: pd.DataFrame, rf: pd.Series,
                       funds: pd.DataFrame) -> pd.DataFrame:
    """Long table: one row per fund per month."""
    long = (monthly.stack(future_stack=True).rename("Return").reset_index()
            .rename(columns={"level_0": "Date", "level_1": "fund_key"}))
    long.columns = ["Date", "fund_key", "Return"]
    long = long.dropna(subset=["Return"])
    long["RF"] = long["Date"].map(rf)
    long = long.dropna(subset=["RF"])
    long["ExcessReturn"] = long["Return"] - long["RF"]

    # benchmark return for the same month (blank for conventional funds)
    bench_map = funds.set_index("fund_key")["benchmark_key"]
    long["benchmark_key"] = long["fund_key"].map(bench_map).astype("Int64")
    bench = long[["Date", "fund_key", "Return"]].rename(
        columns={"fund_key": "benchmark_key", "Return": "BenchmarkReturn"})
    long = long.merge(bench, on=["Date", "benchmark_key"], how="left")
    long["ActiveReturn"] = long["Return"] - long["BenchmarkReturn"]

    # fair-comparison flags
    long["in_pair_period"] = np.where(long["benchmark_key"].isna(), 1,
                                      long["BenchmarkReturn"].notna().astype(int))
    common_start = long.groupby("fund_key")["Date"].min().max()
    long["in_common_period"] = (long["Date"] >= common_start).astype(int)
    return long.sort_values(["fund_key", "Date"]).reset_index(drop=True)


def build_dim_fund(funds: pd.DataFrame, fact: pd.DataFrame,
                   currencies: dict) -> pd.DataFrame:
    d = funds.copy()
    first = fact.groupby("fund_key")["Date"].min()
    last = fact.groupby("fund_key")["Date"].max()
    d["FirstMonth"] = d["fund_key"].map(first)
    d["LastMonth"] = d["fund_key"].map(last)
    d["Months"] = d["fund_key"].map(fact.groupby("fund_key").size())
    d["QuoteCurrency"] = d["ticker"].map(currencies)
    names = d.set_index("fund_key")["short_name"]
    d["BenchmarkName"] = d["benchmark_key"].map(names)
    d["PairName"] = np.where(d["category"] == "ESG",
                             d["short_name"] + " vs " + d["BenchmarkName"].fillna(""),
                             d["short_name"] + " (benchmark)")
    d["CategoryOrder"] = d["category"].map({"ESG": 1, "Conventional": 2})
    return d.drop(columns=["quote_currency_expected"])


def sector_template(funds: pd.DataFrame) -> pd.DataFrame:
    rows = [{"fund_key": k, "ticker": t, "Sector": s, "Weight": None, "AsOf": None}
            for k, t in zip(funds["fund_key"], funds["ticker"]) for s in GICS_SECTORS]
    return pd.DataFrame(rows)


def write_excel_model(path, tables: dict[str, pd.DataFrame]) -> None:
    """One workbook, one sheet per table, each formatted as a named Excel table.

    This is the file to import into the Power BI web service: one upload
    creates one semantic model containing every table, ready for relationships
    and measures. Dates are written as real dates, blanks as empty cells
    (Power BI reads them as null, which the benchmark measures rely on).
    """
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = Workbook()
    wb.remove(wb.active)
    for name, df in tables.items():
        ws = wb.create_sheet(name)
        df = df.copy()
        ws.append(list(df.columns))
        for row in df.itertuples(index=False):
            ws.append([None if (v is None or (not isinstance(v, str) and pd.isna(v)))
                       else (v.to_pydatetime() if isinstance(v, pd.Timestamp)
                             else (v.item() if hasattr(v, "item") else v))
                       for v in row])
        for j, col in enumerate(df.columns, start=1):
            letter = get_column_letter(j)
            if pd.api.types.is_datetime64_any_dtype(df[col]):
                for cell in ws[letter][1:]:
                    cell.number_format = "yyyy-mm-dd"
            ws.column_dimensions[letter].width = max(10, min(40, len(str(col)) + 4))
        ref = f"A1:{get_column_letter(len(df.columns))}{len(df) + 1}"
        table = Table(displayName=name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(table)
        ws.freeze_panes = "A2"
    wb.save(path)


# --------------------------------------------------------------------------
# Python versions of the DAX measures (for checking)
# --------------------------------------------------------------------------
def metrics(fact: pd.DataFrame, period_flag: str | None = "in_pair_period") -> pd.DataFrame:
    """Headline metrics per fund, using the same formulas as measures.dax."""
    f = fact if period_flag is None else fact[fact[period_flag] == 1]
    out = {}
    for k, g in f.groupby("fund_key"):
        r, ex = g["Return"], g["ExcessReturn"]
        n = len(g)
        growth = (1 + r).cumprod()
        dd = growth / np.maximum(growth.cummax(), 1.0) - 1  # peak includes the starting £1
        cagr = (1 + r).prod() ** (12 / n) - 1
        downside = np.sqrt((np.minimum(ex, 0) ** 2).mean()) * np.sqrt(12)
        row = {
            "Months": n,
            "CumulativeReturn": (1 + r).prod() - 1,
            "CAGR": cagr,
            "AnnVolatility": r.std(ddof=1) * np.sqrt(12),
            "Sharpe": ex.mean() / ex.std(ddof=1) * np.sqrt(12),
            "Sortino": ex.mean() * 12 / downside if downside > 0 else np.nan,
            "MaxDrawdown": min(dd.min(), 0.0),
            "WinRate": np.nan,
        }
        b = g["BenchmarkReturn"].dropna()
        if len(b) == n:
            a = g["ActiveReturn"]
            row["BenchmarkCAGR"] = (1 + b).prod() ** (12 / n) - 1
            row["CAGRGap"] = row["CAGR"] - row["BenchmarkCAGR"]
            row["TrackingError"] = a.std(ddof=1) * np.sqrt(12)
            row["InformationRatio"] = a.mean() * 12 / row["TrackingError"]
            row["Beta"] = np.cov(r, b, ddof=1)[0, 1] / b.var(ddof=1)
            row["Correlation"] = np.corrcoef(r, b)[0, 1]
            row["WinRate"] = (a > 0).mean()
            bx = b - g["RF"]
            bg = (1 + b).cumprod()
            row["BenchmarkAnnVolatility"] = b.std(ddof=1) * np.sqrt(12)
            row["BenchmarkSharpe"] = bx.mean() / bx.std(ddof=1) * np.sqrt(12)
            row["BenchmarkMaxDrawdown"] = min((bg / np.maximum(bg.cummax(), 1.0) - 1).min(), 0.0)
        out[k] = row
    return pd.DataFrame(out).T.rename_axis("fund_key")


# --------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------
def build(refresh: bool = False, verbose: bool = True,
          end: pd.Timestamp | None = None) -> dict[str, pd.DataFrame]:
    log = cleaning.QualityLog()
    funds = load_fund_list()
    tickers = funds["ticker"].tolist()

    raw = data_loader.download_yahoo(tickers, DOWNLOAD_START, None,
                                     name="multi_fund_prices_raw", refresh=refresh)
    live = data_loader.yahoo_currencies(tickers)
    fx = data_loader.download_fx(refresh)
    sonia = data_loader.download_sonia(refresh)

    currencies, gbp = {}, {}
    for _, f in funds.iterrows():
        t, k = f["ticker"], f["fund_key"]
        ccy = live.get(t) or f["quote_currency_expected"]
        if live.get(t) and live[t] != f["quote_currency_expected"]:
            warnings.warn(f"{t}: Yahoo currency {live[t]!r} differs from funds.csv "
                          f"{f['quote_currency_expected']!r}; using Yahoo's")
        currencies[t] = ccy
        if t not in raw.columns or raw[t].dropna().empty:
            log.add(f"{t}: NO DATA from Yahoo - check the ticker; fund skipped")
            continue
        s = to_gbp(raw[t].rename(t), ccy, fx, log)
        gbp[k] = clean_daily(s, log)
        if ccy != "GBP" and not ccy.startswith("GB"):
            log.add(f"{t}: quoted in {ccy}, converted to GBP at daily FX")

    daily = pd.DataFrame(gbp)
    monthly = cleaning.monthly_returns(daily, end=end)
    rf = cleaning.monthly_rf_from_sonia(sonia)

    fact = build_fact_returns(monthly, rf, funds[funds["fund_key"].isin(gbp)])
    dim_fund = build_dim_fund(funds[funds["fund_key"].isin(gbp)], fact, currencies)
    dim_date = build_dim_date(pd.DatetimeIndex(sorted(fact["Date"].unique())))
    large = cleaning.flag_large_moves(daily, log)

    for _, f in dim_fund.iterrows():
        log.add(f"{f['ticker']}: {f['FirstMonth']:%b %Y} to {f['LastMonth']:%b %Y} "
                f"({f['Months']} months)")
    log.add(f"Common period (all funds): from "
            f"{fact.loc[fact['in_common_period'] == 1, 'Date'].min():%b %Y}")

    PBI_DATA.mkdir(parents=True, exist_ok=True)
    config.DATA_PROCESSED.mkdir(parents=True, exist_ok=True)
    fmt = {"date_format": "%Y-%m-%d", "index": False}
    dim_fund.to_csv(PBI_DATA / "dim_fund.csv", **fmt)
    dim_date.to_csv(PBI_DATA / "dim_date.csv", **fmt)
    build_dim_regime().to_csv(PBI_DATA / "dim_regime.csv", index=False)
    fact.to_csv(PBI_DATA / "fact_returns_monthly.csv", float_format="%.8f", **fmt)
    sector_path = PBI_DATA / "fact_sector_weights.csv"
    if not sector_path.exists():  # never overwrite weights you have typed in
        sector_template(dim_fund).to_csv(sector_path, index=False)

    check = metrics(fact).join(dim_fund.set_index("fund_key")["short_name"])
    check.to_csv(PBI_DATA / "metrics_check.csv", float_format="%.6f")
    daily.rename(columns=dim_fund.set_index("fund_key")["ticker"]).to_csv(
        config.DATA_PROCESSED / "multi_fund_daily_gbp.csv")
    large.to_csv(config.DATA_PROCESSED / "multi_fund_large_moves.csv")
    write_excel_model(PBI_DATA / "ESG_funds_model.xlsx", {
        "dim_fund": dim_fund, "dim_date": dim_date, "dim_regime": build_dim_regime(),
        "fact_returns_monthly": fact,
        "fact_sector_weights": pd.read_csv(sector_path),
    })
    (PBI_DATA / "data_quality.md").write_text("# Data quality log\n\n" + str(log) + "\n")

    if verbose:
        print("Cleaning log\n" + str(log))
        print("\nHeadline metrics (pair period) - your DAX should match these:")
        cols = ["short_name", "Months", "CAGR", "BenchmarkCAGR", "CAGRGap",
                "Sharpe", "BenchmarkSharpe", "MaxDrawdown", "BenchmarkMaxDrawdown"]
        print(check[[c for c in cols if c in check]].to_string(
            float_format=lambda x: f"{x:.4f}"))
    return {"fact": fact, "dim_fund": dim_fund, "dim_date": dim_date,
            "metrics": check, "log": log}


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--refresh", action="store_true")
    build(refresh=ap.parse_args().refresh)
